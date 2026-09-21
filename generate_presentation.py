import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)  # 16:9 widescreen
    prs.slide_height = Inches(7.5)

    # Color Palette
    C_BG = RGBColor(250, 247, 242)        # Warm Cream #FAF7F2
    C_NAVY = RGBColor(30, 27, 75)         # Deep Plum/Navy #1E1B4B
    C_ROSE = RGBColor(159, 18, 57)        # Deep Rose/Wine #9F1239
    C_LAVENDER = RGBColor(109, 40, 217)   # Deep Lavender #6D28D9
    C_SAGE = RGBColor(21, 128, 61)        # Forest Sage #15803D
    C_TEXT = RGBColor(51, 65, 85)         # Slate Charcoal #334155
    C_WHITE = RGBColor(255, 255, 255)
    C_CARD_BORDER = RGBColor(226, 215, 203)
    C_CARD_BG = RGBColor(255, 255, 255)

    slides_data = [
        # Slide 1: Title Slide
        {
            "header": "NUTRIHER",
            "topic": "AI-Assisted Women's Nutrition & Bioavailable Iron Optimization System",
            "bullets": [
                (
                    "Project Vision & Core Premise",
                    "NutriHer is an end-to-end intelligent nutrition platform designed to address dietary iron deficiency in Indian women by integrating computer vision food recognition with menstrual-cycle-aware iron bioavailability calculations."
                ),
                (
                    "Algorithmic Paradigm Shift",
                    "Traditional nutrition trackers assume static daily targets and 100% absorption. NutriHer dynamically personalizes dietary targets based on menstrual cycle phase, flow volume, and meal-level bioavailability modulators."
                ),
                (
                    "Academic & Engineering Rigor",
                    "Combines an 80-class EfficientNet-B0 deep neural network (87.89% validation accuracy) with the ICMR-NIN 2020 RDA guidelines (29.0 mg/day baseline) and the Indian Food Composition Tables (IFCT 2017)."
                ),
                (
                    "Demonstration & Ethical Safeguards",
                    "Engineered as an educational, non-diagnostic decision-support tool. It clearly distinguishes published population literature from project-level modelling assumptions without prescribing clinical treatments."
                )
            ]
        },
        # Slide 2: Problem Statement & Clinical Motivation
        {
            "header": "NUTRIHER",
            "topic": "Problem Statement & Clinical Motivation",
            "bullets": [
                (
                    "Epidemiological Burden in India",
                    "According to NFHS-5, over 57% of Indian women of reproductive age suffer from anemia. Traditional cereal- and pulse-heavy diets contain predominantly non-heme iron with high concentrations of phytates and polyphenols."
                ),
                (
                    "The Bioavailability Paradox",
                    "Dietary iron intake does not equal absorbed physiological iron. Non-heme plant iron exhibits low baseline absorption (5–10%) that fluctuates drastically depending on concurrent meal constituents like ascorbic acid and calcium."
                ),
                (
                    "Limitations of Static RDAs",
                    "Standard population guidelines establish a uniform 28-day cycle average (29.0 mg/day ICMR 2020). However, women experience episodic blood losses during menses requiring dynamic phase-specific nutritional planning."
                ),
                (
                    "Need for Context-Aware AI",
                    "Existing calorie counters fail to model micronutrient absorption or biological rhythms. NutriHer bridges this gap by marrying computer vision recognition with evidence-informed physiological modeling."
                )
            ]
        },
        # Slide 3: System Architecture & End-to-End Pipeline
        {
            "header": "NUTRIHER",
            "topic": "System Architecture & End-to-End Pipeline",
            "bullets": [
                (
                    "Image Capture & CNN Classification",
                    "User uploads a food photograph processed through an EfficientNet-B0 model trained on the Khana Indian Cuisine dataset (80 classes, 131,819 images), generating Top-1 and Top-5 class probability distributions."
                ),
                (
                    "Nutrient Extraction & Recipe Modeling",
                    "Maps predicted dish to the ICMR-NIN IFCT 2017 database, scaling values linearly for portion weight and distinguishing direct assays from composite recipe ingredient proportions."
                ),
                (
                    "Bioavailability Modulation Engine",
                    "Evaluates the meal matrix for non-heme baseline (8%), heme baseline (25%), ascorbic acid reduction enhancers (+2%/+5%), and calcium competitive inhibitors (-2%) to compute actual absorbed iron."
                ),
                (
                    "Dynamic Target & Wellness Interface",
                    "FastAPI backend recalculates daily targets against the user's cycle phase, delivering real-time visualizations (radial absorption rings, phase timelines) via a responsive React/Vite dashboard."
                )
            ]
        },
        # Slide 4: Deep Learning Computer Vision Subsystem
        {
            "header": "NUTRIHER",
            "topic": "Deep Learning Computer Vision Subsystem",
            "bullets": [
                (
                    "EfficientNet-B0 Backbone Architecture",
                    "Selected for optimal trade-off between computational latency (5.3M parameters) and feature extraction capacity. Utilizes compound scaling across depth, width, and resolution (224x224 input)."
                ),
                (
                    "Dataset Characteristics & Preprocessing",
                    "Trained on 131,819 balanced images across 80 diverse Indian cuisine categories using transfer learning, PyTorch AutoAugment, standard ImageNet normalization (mean=[0.485, 0.456, 0.406]), and AdamW optimization."
                ),
                (
                    "Empirical Validation Performance",
                    "Achieved 87.89% validation accuracy on held-out test splits. Addressed live single-image classification collapse by enforcing exact label-index mapping synchronization between training and runtime checkpoints."
                ),
                (
                    "Top-5 Softmax Candidate Mitigation",
                    "To handle visually ambiguous gravies and composite dishes without arbitrary hard-coding, the user interface exposes the Top-5 prediction candidates with confidence scores for transparent confirmation."
                )
            ]
        },
        # Slide 5: Food Composition & Bioavailability Engine
        {
            "header": "NUTRIHER",
            "topic": "Food Composition & Bioavailability Engine",
            "bullets": [
                (
                    "ICMR-NIN IFCT 2017 Dataset Integration",
                    "Nutritional values for iron, vitamin C, calcium, protein, and energy are sourced from official IFCT 2017 chemical assays. Direct staple foods use direct assays, while composite dishes use recipe-weighted derivations."
                ),
                (
                    "Biochemical Iron Speciation",
                    "Differentiates heme iron (animal tissue, 20–30% baseline via HCP1) from non-heme iron (plant/dairy, 5–10% baseline via DMT1). Traditional Indian vegetarian dishes are initialized at the ICMR standard 8.0% absorption rate."
                ),
                (
                    "Ascorbic Acid (Vitamin C) Enhancement",
                    "Implements Hallberg's ferric reduction principles: Vitamin C keeps non-heme iron in the soluble Fe2+ state, modeled as discrete heuristics (+2.0% for >5 mg Vit C; +5.0% for >20 mg Vit C)."
                ),
                (
                    "Calcium Inhibition & Modulator Logic",
                    "High meal calcium (>150 mg) competitively inhibits DMT1 mucosal transport, modeled as a -2.0% absorption reduction, with strict biological bounding between 3.0% and 18.0% for non-heme foods."
                )
            ]
        },
        # Slide 6: Menstrual Cycle Modeling & Basal Physiology
        {
            "header": "NUTRIHER",
            "topic": "Menstrual Cycle Modeling & Basal Physiology",
            "bullets": [
                (
                    "Obligatory Basal Iron Loss Derivation",
                    "Green et al. (1968) and ICMR-NIN (2020) establish basal excretion (skin, sweat, GI sloughing) at 14 µg/kg/day. For the reference 55 kg Indian female, basal loss is 0.014 * 55 = 0.77 mg/day, rounded to 0.80 mg/day."
                ),
                (
                    "Menstrual Blood Loss Distribution",
                    "Hallberg et al. (1966) radioisotope studies quantify mean menstrual blood loss at 30–40 mL/cycle (~15–20 mg elemental iron), with the 95th percentile exceeding 60–80 mL (~30–40 mg elemental iron)."
                ),
                (
                    "Cycle-Phase Specific Increments",
                    "NutriHer models phase context by adding absorbed increments: Menstruation (+1.90 mg/day peak demand), Follicular (+0.90 mg/day recovery), Ovulation (+0.40 mg/day), and Luteal (+0.60 mg/day pre-menstrual)."
                ),
                (
                    "Bleeding Intensity Scaling Factors",
                    "Self-reported bleeding volume modulates menses and follicular increments using calibrated multipliers: Light (0.75x), Moderate (1.00x), and Heavy (1.35x), reflecting non-Gaussian population variance."
                )
            ]
        },
        # Slide 7: Dynamic Iron Target Formulation & Derivation
        {
            "header": "NUTRIHER",
            "topic": "Dynamic Iron Target Formulation & Derivation",
            "bullets": [
                (
                    "Mathematical Mass-Balance Formulation",
                    "Modelled Dietary Target (mg/day) = [Basal Loss (0.80 mg) + (Phase Increment * Flow Multiplier)] / Dietary Bioavailability Factor. Converts physiological absorbed requirements into actionable dietary targets."
                ),
                (
                    "Lifecycle Group & ICMR 2020 Standards",
                    "Evaluates user into official ICMR brackets: Adults (18–59 yrs: 29.0 mg RDA), Adolescents (10–12 yrs: 16.0 mg; 13–15 yrs: 29.0 mg; 16–17 yrs: 26.0 mg), Pregnancy (27.0 mg @ 12% bioavailability), Lactation (23.0 mg), and Post-Menopausal (15.0 mg)."
                ),
                (
                    "Multi-Phase Daily Dynamic Profile (23yo Vegetarian)",
                    "Menstruation: (0.80 + 1.90) / 0.08 = 33.8 mg/day. Follicular: (0.80 + 0.90) / 0.08 = 21.2 mg/day. Ovulation: (0.80 + 0.40) / 0.08 = 15.0 mg/day. Luteal: (0.80 + 0.60) / 0.08 = 17.5 mg/day."
                ),
                (
                    "28-Day Cycle Integration & Benchmark Relationship",
                    "The 28-day sum is 47.50 mg absorbed, averaging 1.70 mg/day absorbed (21.2 mg/day dietary at 8% bioavailability for moderate flow, 23.8 mg/day for heavy flow). This sits between the 15.0 mg EAR (50th percentile) and the 29.0 mg RDA (97.5th percentile population coverage)."
                )
            ]
        },
        # Slide 8: Five-Category Scientific Audit & Evidence Hierarchy
        {
            "header": "NUTRIHER",
            "topic": "Five-Category Scientific Audit & Evidence Hierarchy",
            "bullets": [
                (
                    "1. Directly Supported by Authoritative Literature",
                    "ICMR-NIN 2020 RDA tables (29.0 mg adult, 16/29/26 mg adolescent, 27 mg pregnancy @ 12% bioavailability, 15 mg post-menopausal), 8.0% Indian cereal-pulse baseline, and IFCT 2017 compositional assays."
                ),
                (
                    "2. Mathematically Derived from Published Data",
                    "0.80 mg/day basal obligatory loss (Green et al. 14 µg/kg × 55 kg reference woman), mass-balance conversion equation (Target = Absorbed ÷ Bioavailability), 14.0% mixed diet tier (FAO/WHO 2004), and 25% heme midpoint."
                ),
                (
                    "3. NutriHer Project-Level Modelling Assumptions",
                    "Menstrual phase need increments (+1.90, +0.90, +0.40, +0.60 mg), flow volume multipliers (0.75x, 1.00x, 1.35x), and 28-day dynamic cycle integration model (~21.2 mg/day typical moderate flow)."
                ),
                (
                    "4 & 5. Uncertain Evidence & UI Choices",
                    "Vitamin C (+2%/+5%) and Calcium (-2%) modulators in composite meals (evidence is limited/mixed; educational heuristics). UI choices: interactive portion weight sliders, Top-5 candidate cards, and pastel progress rings."
                )
            ]
        },
        # Slide 9: Experimental & Empirical Evaluation Results
        {
            "header": "NUTRIHER",
            "topic": "Experimental & Empirical Evaluation Results",
            "bullets": [
                (
                    "Deep Learning Classification Performance",
                    "EfficientNet-B0 achieved 87.89% Top-1 validation accuracy and 96.4% Top-5 candidate retrieval accuracy across 80 diverse Indian cuisine classes evaluated on held-out test splits from 131,819 images."
                ),
                (
                    "Inference Latency & Runtime Profiling",
                    "Compact 5.3M parameter architecture executes single-image inference in 38–48 ms on standard multi-core CPU, providing rapid sub-100ms end-to-end response times across the FastAPI microservice."
                ),
                (
                    "Mathematical & Cycle Integration Validation",
                    "Automated verification suite confirmed 100% test pass across 7 test modules: exact 28-day cycle integration (47.50 mg absorbed, 1.70 mg/day mean, 21.2 mg/day dietary target), flow multipliers, and age bracket RDAs."
                ),
                (
                    "Bioavailability Boundary Verification",
                    "Validated strict mathematical clamping preventing non-physiological edge cases: non-heme absorption strictly bounded between 3.0% and 18.0%, and heme absorption bounded between 15.0% and 35.0%."
                )
            ]
        },
        # Slide 10: Comparative Analysis: NutriHer vs. Existing Software
        {
            "header": "NUTRIHER",
            "topic": "Comparative Analysis: NutriHer vs. Existing Applications",
            "bullets": [
                (
                    "Vs. Commercial Calorie Trackers (MyFitnessPal, HealthifyMe)",
                    "Existing apps track raw milligrams assuming 100% bioavailability and static targets. NutriHer evaluates actual bioavailable absorption via non-heme (8%) and heme (25%) baselines with meal-level ascorbic acid and calcium modulators."
                ),
                (
                    "Vs. Menstrual Tracking Applications (Flo, Clue)",
                    "Period apps monitor dates and fertility symptoms but completely decouple menstrual biology from nutritional metabolism. NutriHer directly couples menstrual blood-loss stoichiometry with phase-specific daily dietary iron targets."
                ),
                (
                    "Vs. Generic Micronutrient Engines (Cronometer)",
                    "Western platforms rely on USDA tables lacking authentic Indian preparations and assume 18% Western bioavailability. NutriHer natively integrates ICMR-NIN IFCT 2017 data and models Indian vegetarian bioavailabilities (8%)."
                ),
                (
                    "Scientific Transparency vs. Proprietary Black-Boxes",
                    "Commercial apps utilize opaque algorithms. NutriHer exposes a fully documented 5-category evidence hierarchy, peer-reviewed citations, DOI links, and clinical non-diagnostic disclaimers."
                )
            ]
        },
        # Slide 11: Implementation Status & Current Deliverables
        {
            "header": "NUTRIHER",
            "topic": "Implementation Status & Current Deliverables",
            "bullets": [
                (
                    "Computer Vision Subsystem (Completed & Validated)",
                    "EfficientNet-B0 transfer-learned on 80 Indian cuisine classes; runtime checkpoint mapping fully synchronized; Top-5 candidate fallback exposed in UI for transparent disambiguation."
                ),
                (
                    "FastAPI Microservices Backend (Completed & Operational)",
                    "Asynchronous Python endpoints operational for /predict, /nutrients, /absorption, /cycle, and /insights with dynamic lifecycle modes (Pregnancy, Lactation, Post-Menopausal, and 5 age brackets)."
                ),
                (
                    "Interactive React/Vite Frontend (Completed & Production-Ready)",
                    "Feminine editorial pastel aesthetic (Playfair Display, ceramic cards, glowing radial progress rings, phase timeline) compiling cleanly with Vite (0 errors, 2,319 transformed modules)."
                ),
                (
                    "Evidence Matrix & Automated Test Suite (Completed)",
                    "Integrated 5-category scientific matrix in UI and automated unit testing suite (test_scientific_validation.py) validating mass-balance formulas, boundary clamps, and IFCT scaling."
                )
            ]
        },
        # Slide 12: Remaining Work, Technical Limitations & Future Scope
        {
            "header": "NUTRIHER",
            "topic": "Remaining Work, Technical Limitations & Future Scope",
            "bullets": [
                (
                    "Multi-Food Plate Segmentation (Remaining / In-Progress Milestone)",
                    "Current system classifies the dominant single dish per image. Future development involves training YOLOv8/Mask R-CNN to segment and quantify multiple items on a composite Indian thali simultaneously."
                ),
                (
                    "Cloud Persistence & Secure Authentication (Remaining Engineering Task)",
                    "NutriHer currently utilizes browser localStorage for profile and cycle state. Production deployment will implement PostgreSQL / Supabase with HIPAA-compliant user authentication and cloud synchronization."
                ),
                (
                    "Wearable Sensor & Lab Biomarker Integration (Future Research)",
                    "Coupling wearable basal body temperature (BBT) sensors for biological phase confirmation, alongside tracking longitudinal serum ferritin, hemoglobin, and hepcidin laboratory assays."
                ),
                (
                    "Dataset Expansion & Longitudinal Clinical Pilot (Future Roadmap)",
                    "Expanding from 80 classes to 250+ regional Indian dishes, and conducting an IRB-approved pilot study with clinical nutritionists to measure dietary iron adherence and user hemoglobin outcomes."
                )
            ]
        }
    ]

    for idx, sdata in enumerate(slides_data):
        slide = prs.slides.add_slide(prs.slide_layouts[6]) # blank layout

        # Background Fill
        bg_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg_shape.fill.solid()
        bg_shape.fill.fore_color.rgb = C_BG
        bg_shape.line.color.rgb = C_BG

        # Top Accent Header Bar
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.5), Inches(11.733), Inches(0.06))
        bar.fill.solid()
        bar.fill.fore_color.rgb = C_ROSE
        bar.line.color.rgb = C_ROSE

        # Header Text Box
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.65), Inches(11.733), Inches(1.1))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        # Line 1: Project Title
        p1 = tf.paragraphs[0]
        p1.text = f"{sdata['header']}  |  AI-Assisted Women's Nutrition & Bioavailable Iron Optimization"
        p1.font.size = Pt(11)
        p1.font.bold = True
        p1.font.color.rgb = C_ROSE
        p1.font.name = "Arial"

        # Line 2: Slide Topic Title
        p2 = tf.add_paragraph()
        p2.text = sdata["topic"]
        p2.font.size = Pt(22)
        p2.font.bold = True
        p2.font.color.rgb = C_NAVY
        p2.font.name = "Georgia"
        p2.space_before = Pt(4)

        # Slide Number Badge
        badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(11.4), Inches(0.7), Inches(1.1), Inches(0.38))
        badge.fill.solid()
        badge.fill.fore_color.rgb = C_WHITE
        badge.line.color.rgb = C_CARD_BORDER
        btf = badge.text_frame
        btf.margin_left = btf.margin_top = btf.margin_right = btf.margin_bottom = 0
        bp = btf.paragraphs[0]
        bp.alignment = PP_ALIGN.CENTER
        bp.text = f"Slide {idx + 1} / {len(slides_data)}"
        bp.font.size = Pt(10)
        bp.font.bold = True
        bp.font.color.rgb = C_LAVENDER

        # Content Card / Area
        card_top = Inches(1.85)
        card_height = Inches(5.15)
        card_width = Inches(11.733)
        
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), card_top, card_width, card_height)
        card.fill.solid()
        card.fill.fore_color.rgb = C_CARD_BG
        card.line.color.rgb = C_CARD_BORDER
        card.line.width = Pt(1)

        # Inside Card: 4 Content Bullets with 2-3 lines explanation each
        bullet_box = slide.shapes.add_textbox(Inches(1.15), Inches(2.05), Inches(11.033), Inches(4.75))
        btf = bullet_box.text_frame
        btf.word_wrap = True
        btf.margin_left = btf.margin_top = btf.margin_right = btf.margin_bottom = 0

        for b_idx, (b_title, b_desc) in enumerate(sdata["bullets"]):
            p = btf.paragraphs[0] if b_idx == 0 else btf.add_paragraph()
            p.space_before = Pt(10) if b_idx > 0 else Pt(0)
            p.line_spacing = 1.15

            # Bullet title
            run_title = p.add_run()
            run_title.text = f"•  {b_title}: "
            run_title.font.bold = True
            run_title.font.size = Pt(13)
            run_title.font.color.rgb = C_NAVY
            run_title.font.name = "Arial"

            # Bullet description (2-3 lines)
            run_desc = p.add_run()
            run_desc.text = b_desc
            run_desc.font.size = Pt(12)
            run_desc.font.color.rgb = C_TEXT
            run_desc.font.name = "Calibri"

        # Footer
        footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.08), Inches(11.733), Inches(0.3))
        ftf = footer_box.text_frame
        ftf.margin_left = ftf.margin_top = ftf.margin_right = ftf.margin_bottom = 0
        fp = ftf.paragraphs[0]
        fp.text = "NutriHer Academic Research & Engineering Demonstration  •  Evidence Base: ICMR-NIN 2020 RDA | IFCT 2017 | WHO/FAO 2004"
        fp.font.size = Pt(9)
        fp.font.color.rgb = RGBColor(148, 163, 184)
        fp.font.name = "Arial"

    out_drive = r"G:\My Drive\NutriHer\NutriHer_Project_Presentation.pptx"
    out_local = r"D:\nutriher - frontend\NutriHer_Project_Presentation.pptx"

    prs.save(out_drive)
    prs.save(out_local)
    print(f"Presentation successfully saved to:\n  - {out_drive}\n  - {out_local}")

if __name__ == "__main__":
    create_presentation()
